"""Render user-owned workspace text into downloadable office documents.

The model produces bounded UTF-8 text, never binary, macros, formulas or code.
Rendering is deterministic, uses in-memory buffers and executes no user content.
"""
from __future__ import annotations

import csv
from html import escape
from io import BytesIO, StringIO
from pathlib import PurePath

from .file_workspace import WorkspaceError

MAX_OUTPUT_BYTES = 4_000_000
MAX_SPREADSHEET_ROWS = 2000
MAX_SPREADSHEET_COLUMNS = 50


def _markdown_line(line: str) -> tuple[str, str]:
    stripped = line.strip()
    if stripped.startswith("#"):
        level = len(stripped) - len(stripped.lstrip("#"))
        if 1 <= level <= 3 and stripped[level:level + 1] == " ":
            return ("heading" + str(level), stripped[level:].strip())
    if stripped.startswith(("- ", "* ")):
        return ("bullet", stripped[2:].strip())
    return ("paragraph", line)


def _render_pdf(content: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    buffer = BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="IAChatBody", parent=styles["BodyText"],
                              fontName="Helvetica", fontSize=10, leading=14, spaceAfter=7))
    styles.add(ParagraphStyle(name="IAChatBullet", parent=styles["IAChatBody"],
                              leftIndent=16, firstLineIndent=-9))
    story = []
    for line in content.splitlines():
        kind, text = _markdown_line(line)
        if not text.strip():
            story.append(Spacer(1, 7))
        elif kind.startswith("heading"):
            story.append(Paragraph(escape(text), styles["Heading" + kind[-1]]))
        elif kind == "bullet":
            story.append(Paragraph("&#8226; " + escape(text), styles["IAChatBullet"]))
        else:
            story.append(Paragraph(escape(text), styles["IAChatBody"]))
    if not story:
        story.append(Paragraph(" ", styles["IAChatBody"]))
    document = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=48,
                                 rightMargin=48, topMargin=48, bottomMargin=48,
                                 title="Documento de IAchat")
    document.build(story)
    return buffer.getvalue()


def _render_docx(content: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for line in content.splitlines():
        kind, text = _markdown_line(line)
        if kind.startswith("heading"):
            document.add_heading(text, level=int(kind[-1]))
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_xlsx(content: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    sheet.freeze_panes = "A2"
    widths: list[int] = []
    try:
        reader = csv.reader(StringIO(content, newline=""), strict=True)
        for row_number, row in enumerate(reader, 1):
            if row_number > MAX_SPREADSHEET_ROWS or len(row) > MAX_SPREADSHEET_COLUMNS:
                raise WorkspaceError("spreadsheet_limit_reached")
            for column_number, value in enumerate(row, 1):
                if len(value) > 32767:
                    raise WorkspaceError("spreadsheet_cell_too_large")
                cell = sheet.cell(row=row_number, column=column_number, value=value)
                # Treat untrusted CSV as text. In particular, never interpret
                # strings starting with '=' as executable spreadsheet formulas.
                cell.data_type = "s"
                if row_number == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="25466A")
                    cell.alignment = Alignment(wrap_text=True)
                if len(widths) < column_number:
                    widths.append(10)
                widths[column_number - 1] = min(40, max(widths[column_number - 1], len(value) + 2))
    except (csv.Error, UnicodeError) as exc:
        raise WorkspaceError("invalid_spreadsheet_csv") from exc
    if sheet.max_row > 1 and sheet.max_column:
        sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_download(name: str, content: str) -> tuple[bytes, str]:
    """Return bytes and MIME type. All input comes from an authenticated workspace."""
    extension = PurePath(name).suffix.lower()
    if extension == ".pdf":
        payload, media_type = _render_pdf(content), "application/pdf"
    elif extension == ".docx":
        payload, media_type = (_render_docx(content),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    elif extension == ".xlsx":
        payload, media_type = (_render_xlsx(content),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        payload, media_type = content.encode("utf-8"), "application/octet-stream"
    if len(payload) > MAX_OUTPUT_BYTES:
        raise WorkspaceError("rendered_file_too_large")
    return payload, media_type
