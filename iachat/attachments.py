"""Bounded local extraction of user-provided documents. No external network calls."""
import base64
import binascii
import io
from pathlib import PurePath

MAX_FILE_BYTES = 2_000_000
MAX_TOTAL_CHARS = 18_000

class AttachmentError(ValueError):
    pass

def extract(attachments):
    remaining = MAX_TOTAL_CHARS
    sections = []
    names = []
    if len(attachments) > 3:
        raise AttachmentError("too_many_attachments")
    for item in attachments:
        name = PurePath(item.name.replace("\\", "/")).name[:180]
        ext = name.lower().rsplit(".", 1)[-1]
        if ext not in ("txt", "md", "pdf", "docx", "xlsx"):
            raise AttachmentError("unsupported_attachment_type")
        if len(item.content_base64) > 3_000_000:
            raise AttachmentError("attachment_too_large")
        try:
            raw = base64.b64decode(item.content_base64, validate=True)
        except (ValueError, binascii.Error):
            raise AttachmentError("invalid_attachment_encoding")
        if len(raw) > MAX_FILE_BYTES:
            raise AttachmentError("attachment_too_large")
        if ext in ("docx", "xlsx"):
            # Office files are ZIP containers. Limit decompressed size before
            # asking the document libraries to parse attacker-controlled XML.
            try:
                from zipfile import ZipFile
                with ZipFile(io.BytesIO(raw)) as archive:
                    members = archive.infolist()
                    if len(members) > 1000 or sum(m.file_size for m in members) > 15_000_000:
                        raise AttachmentError("office_document_too_large")
                    if any(m.flag_bits & 1 for m in members):
                        raise AttachmentError("encrypted_office_document")
                if ext == "docx":
                    from docx import Document
                    doc = Document(io.BytesIO(raw))
                    paragraphs = [p.text for p in doc.paragraphs[:400] if p.text.strip()]
                    for table in doc.tables[:20]:
                        for row in table.rows[:100]:
                            paragraphs.append(" | ".join(cell.text for cell in row.cells[:30]))
                    content = "\\n".join(paragraphs)
                else:
                    from openpyxl import load_workbook
                    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                    try:
                        lines = []
                        for sheet in book.worksheets[:5]:
                            lines.append("[HOJA: " + sheet.title + "]")
                            for row in sheet.iter_rows(min_row=1, max_row=200, max_col=30, values_only=True):
                                if any(value is not None for value in row):
                                    lines.append(" | ".join("" if value is None else str(value)[:500] for value in row))
                                if sum(len(line) for line in lines) > MAX_TOTAL_CHARS:
                                    break
                        content = "\\n".join(lines)
                    finally:
                        book.close()
            except AttachmentError:
                raise
            except Exception as exc:
                raise AttachmentError("office_document_unreadable") from exc
        elif ext == "pdf":
            try:
                from pypdf import PdfReader
                pdf = PdfReader(io.BytesIO(raw), strict=False)
                if len(pdf.pages) > 40:
                    raise AttachmentError("pdf_too_many_pages")
                content = "\n".join(page.extract_text() or "" for page in pdf.pages)
            except AttachmentError:
                raise
            except Exception:
                raise AttachmentError("pdf_text_unreadable")
        else:
            try:
                content = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise AttachmentError("text_file_not_utf8")
        content = content.strip()
        if not content:
            raise AttachmentError("attachment_contains_no_text")
        names.append(name)
        sections.append("[DOCUMENTO: " + name + "]\n" + content[:remaining] + "\n[/DOCUMENTO]")
        remaining -= min(remaining, len(content))
        if remaining <= 0:
            break
    return "\n\n".join(sections), names


def extract_images(attachments):
    """Validate still images and return ephemeral data URLs; never persist bytes."""
    images = []
    names = []
    formats = {
        "jpg": ("jpeg", lambda b: b.startswith(bytes.fromhex("ffd8ff"))),
        "jpeg": ("jpeg", lambda b: b.startswith(bytes.fromhex("ffd8ff"))),
        "png": ("png", lambda b: b.startswith(bytes.fromhex("89504e470d0a1a0a"))),
        "webp": ("webp", lambda b: b.startswith(b"RIFF") and b[8:12] == b"WEBP"),
    }
    if len(attachments) > 5:
        raise AttachmentError("too_many_attachments")
    for item in attachments:
        name = PurePath(item.name.replace("\\", "/")).name[:180]
        ext = name.lower().rsplit(".", 1)[-1]
        if ext not in formats:
            raise AttachmentError("unsupported_image_type")
        if len(item.content_base64) > 2_750_000:
            raise AttachmentError("image_too_large")
        try:
            raw = base64.b64decode(item.content_base64, validate=True)
        except (ValueError, binascii.Error):
            raise AttachmentError("invalid_image_encoding")
        if not raw or len(raw) > MAX_FILE_BYTES:
            raise AttachmentError("image_too_large")
        mime, signature_ok = formats[ext]
        if not signature_ok(raw):
            raise AttachmentError("image_signature_invalid")
        images.append("data:image/" + mime + ";base64," + item.content_base64)
        names.append(name)
    return images, names
