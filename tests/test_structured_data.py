"""Real authenticated HTTP tests for owner-isolated structured data and Excel."""
import json
from io import BytesIO
from openpyxl import load_workbook
from fastapi.testclient import TestClient
import iachat.app as app_module
from iachat.store import ConversationStore
from iachat.file_workspace import FileWorkspace
from iachat.structured_data import StructuredStore, DataError


class Planner:
    def answer(self, messages, images=None):
        request = json.loads(messages[-1]["content"])
        text = request["request"].casefold()
        collections = request["collections"]
        if "creá planilla gastos" in text:
            return json.dumps({"action":"create","name":"gastos","fields":[
                {"name":"Fecha","type":"date"},{"name":"Concepto","type":"text"},
                {"name":"Categoría","type":"text"},{"name":"Importe","type":"number"}]})
        if "creá colección inventario" in text:
            return json.dumps({"action":"create","name":"inventario","fields":[
                {"name":"Producto","type":"text"},{"name":"Cantidad","type":"number"},
                {"name":"Ubicación","type":"text"}]})
        collection = next((c for c in collections if c["name"] in ("inventario","gastos")
                           and (("inventario" in text or "tornillos" in text) == (c["name"]=="inventario"))),collections[0])
        cid = collection["id"]
        if "nafta" in text:
            return json.dumps({"action":"add","collection_id":cid,"values":{
                "Fecha":request["today_argentina"],"Concepto":"nafta",
                "Categoría":"transporte","Importe":40000}})
        if "tornillos" in text:
            return json.dumps({"action":"add","collection_id":cid,"values":{
                "Producto":"tornillos","Cantidad":12,"Ubicación":"depósito"}})
        if "cuánto" in text:
            return json.dumps({"action":"query","collection_id":cid,"aggregate":"sum",
                               "field":"Importe","filters":{"Categoría":"transporte"}})
        if "corregí" in text:
            return json.dumps({"action":"update","collection_id":cid,"values":{"Importe":45000}})
        if "descargá" in text:
            return json.dumps({"action":"export","collection_id":cid})
        return '{"action":"none"}'


def test_persistent_collections_two_users_and_sessions(monkeypatch,tmp_path):
    db = str(tmp_path/"data.db")
    monkeypatch.setattr(app_module,"store",ConversationStore(db))
    monkeypatch.setattr(app_module,"file_workspace",FileWorkspace(db))
    monkeypatch.setattr(app_module,"structured_store",StructuredStore(db))
    monkeypatch.setattr(app_module,"get_backend",lambda:Planner())
    monkeypatch.setenv("IACHAT_AUTH_SECRET","test-secret")
    assert app_module.store.create_user("alice","password123")
    assert app_module.store.create_user("bob","password123")
    with TestClient(app_module.app,base_url="https://testserver") as a, TestClient(app_module.app,base_url="https://testserver") as b:
        assert a.post("/auth/login",json={"username":"alice","password":"password123"}).status_code==200
        assert b.post("/auth/login",json={"username":"bob","password":"password123"}).status_code==200
        def chat(client,text):
            response=client.post("/api/chat",json={"text":text})
            assert response.status_code==200,response.text
            return response.json()
        created=chat(a,"creá planilla gastos fecha concepto categoría importe")
        assert "colección persistente" in created["answer"]
        assert "gastos" in created["answer"]
        assert app_module.structured_store.collections("bob")==[]
        assert "Registrado" in chat(a,"anotá 40000 nafta hoy")["answer"]
        assert "40000" in chat(a,"¿cuánto gasté en transporte?")["answer"]
        assert "45000" in chat(a,"corregí el importe a 45000")["answer"]
        # Same user, another browser session and a fresh conversation.
        with TestClient(app_module.app,base_url="https://testserver") as second:
            second.post("/auth/login",json={"username":"alice","password":"password123"})
            assert "45000" in chat(second,"¿cuánto gasté en transporte?")["answer"]
            exported=chat(second,"descargá Excel")
            info=exported["generated_file"]
            assert info["name"].endswith(".xlsx")
            assert "/api/files/"+info["id"]+"/download" in exported["answer"]
            downloaded=second.get("/api/files/"+info["id"]+"/download")
            assert downloaded.status_code==200
            sheet=load_workbook(BytesIO(downloaded.content)).active
            assert [sheet.cell(1,i).value for i in range(1,5)]==["Fecha","Concepto","Categoría","Importe"]
            assert sheet["D2"].value==45000
            assert sheet["D3"].value=="=SUM(D2:D2)"
            assert sheet["D3"].data_type=="f"
            assert b.get("/api/files/"+info["id"]+"/download").status_code==404
        assert "colección persistente" in chat(b,"creá colección inventario producto cantidad ubicación")["answer"]
        assert "Registrado" in chat(b,"anotá 12 tornillos en depósito")["answer"]
        assert len(app_module.structured_store.records("bob",app_module.structured_store.collections("bob")[0]["id"]))==1
        assert len(app_module.structured_store.records("alice",app_module.structured_store.collections("alice")[0]["id"]))==1
        assert app_module.structured_store.collections("alice")[0]["name"]=="gastos"


def test_ambiguous_edit_and_cross_owner_fail_closed(tmp_path):
    data=StructuredStore(str(tmp_path/"data.db"))
    coll=data.create("alice","vehículos",[{"name":"Patente","type":"text"},{"name":"Kilómetros","type":"number"}])
    cid=coll["id"]
    data.add("alice",cid,{"Patente":"ABC123","Kilómetros":120})
    data.add("alice",cid,{"Patente":"DEF456","Kilómetros":240})
    from iachat.structured_data import handle
    class Ambiguous:
        def answer(self,messages):
            return json.dumps({"action":"update","collection_id":cid,"values":{"Kilómetros":300}})
    files=FileWorkspace(str(tmp_path/"data.db"))
    result=handle("corregí kilómetros a 300","alice",Ambiguous(),data,files)
    assert "No modifiqué" in result["answer"]
    assert sorted(r["values"]["Kilómetros"] for r in data.records("alice",cid))==[120,240]
    assert data.collections("bob")==[]
    try:
        data.records("bob",cid)
        assert False
    except DataError:
        pass


def test_arbitrary_numeric_field_exports_as_number_and_formula(tmp_path):
    db=str(tmp_path/"custom.db")
    data=StructuredStore(db)
    files=FileWorkspace(db)
    collection=data.create("alice","flota",[
        {"name":"Unidad","type":"text"},{"name":"Kilómetros","type":"number"}])
    data.add("alice",collection["id"],{"Unidad":"Fiorino","Kilómetros":21345})
    saved=data.export("alice",collection["id"],files)
    from iachat.document_render import render_download
    row=files.get("alice",saved["id"])
    payload,media=render_download(row["name"],row["content"])
    sheet=load_workbook(BytesIO(payload)).active
    assert sheet["B2"].value==21345 and sheet["B2"].data_type=="n"
    assert sheet["B3"].value=="=SUM(B2:B2)" and sheet["B3"].data_type=="f"
    assert files.get("bob",saved["id"]) is None
