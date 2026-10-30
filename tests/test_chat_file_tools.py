"""The regular chat endpoint can create/edit real private downloadable files."""
import json
from io import BytesIO

from openpyxl import load_workbook

from fastapi.testclient import TestClient

import iachat.app as app_module
from iachat.file_workspace import FileWorkspace
from iachat.store import ConversationStore
from iachat.tool_planner import FileAction, is_file_action_candidate, select_file_action


class FakeBackend:
    def __init__(self):
        self.file_id = None

    def answer(self, messages, images=None):
        policy = messages[0]["content"] if messages[0]["role"] == "system" else ""
        request = messages[-1]["content"]
        if "selector de herramientas" in policy:
            task = json.loads(request)["request"]
            if task.startswith("Creá"):
                return '{"action":"create"}'
            if task.startswith("Actualizá"):
                return json.dumps({"action": "edit", "file_id": self.file_id})
            return '{"action":"answer"}'
        if "generador y editor" in policy:
            if self.file_id and "DOCUMENTO EXISTENTE" in request:
                return json.dumps({"name": "malicious-change.pdf", "content": "# Informe\n\nVersión corregida"})
            return json.dumps({"name": "informe.pdf", "content": "# Informe\n\nContenido de prueba"})
        return "Una respuesta normal."


def test_file_prefilter_only_selects_classifier_candidates():
    assert is_file_action_candidate("Creá un informe PDF descargable")
    assert is_file_action_candidate("Creame una planilla Excel con total automático para descargar")
    assert is_file_action_candidate("Actualizá el informe que acabás de crear")
    assert not is_file_action_candidate("¿Qué es un PDF?")
    assert not is_file_action_candidate("¿Cuándo juega River?")
    assert not is_file_action_candidate("Traducí este texto al inglés")
    assert not is_file_action_candidate("Mi archivo está en una carpeta")


def test_selector_rejects_unknown_file_ids_and_ambiguous_output():
    class Backend:
        def __init__(self, output):
            self.output = output
        def answer(self, messages):
            return self.output
    files = [{"id": "owned", "name": "informe.pdf"}]
    assert select_file_action("Editá el archivo", files, Backend('{"action":"edit","file_id":"foreign"}')) == FileAction()
    assert select_file_action("Editá el archivo", files, Backend('{"action":"edit","file_id":"owned"}')) == FileAction("edit", "owned")
    assert select_file_action("Hola", files, Backend('{"action":"create","extra":true}')) == FileAction()
    assert select_file_action("Hola", files, Backend("texto inesperado")) == FileAction()
    assert select_file_action("Hola", files, Backend("ERROR")) == FileAction()


def test_authenticated_chat_creates_edits_and_downloads_pdf(monkeypatch, tmp_path):
    db = str(tmp_path / "chat-tools.sqlite")
    monkeypatch.setattr(app_module, "store", ConversationStore(db))
    monkeypatch.setattr(app_module, "file_workspace", FileWorkspace(db))
    monkeypatch.setenv("IACHAT_AUTH_SECRET", "only-for-isolated-tests")
    assert app_module.store.create_user("alice", "password123", admin=True)
    assert app_module.store.create_user("bob", "password123")
    backend = FakeBackend()
    monkeypatch.setattr(app_module, "get_backend", lambda: backend)
    monkeypatch.setattr("iachat.auto_memory.learn", lambda *args: None)
    with TestClient(app_module.app, base_url="https://testserver") as client:
        assert client.post("/auth/login", json={"username": "alice", "password": "password123"}).status_code == 200
        created = client.post("/api/chat", json={"text": "Creá un informe PDF descargable sobre una prueba."})
        assert created.status_code == 200, created.text
        payload = created.json()
        file_id = payload["generated_file"]["id"]
        backend.file_id = file_id
        assert payload["generated_file"]["name"] == "informe.pdf"
        assert f"/api/files/{file_id}/download" in payload["answer"]
        pdf = client.get(f"/api/files/{file_id}/download")
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")
        edited = client.post("/api/chat", json={
            "conversation_id": payload["conversation_id"],
            "text": "Actualizá el informe que acabás de crear.",
        })
        assert edited.status_code == 200, edited.text
        assert edited.json()["generated_file"]["id"] == file_id
        assert "actualizado" in edited.json()["answer"].lower()
        assert "Versión corregida" in app_module.file_workspace.get("alice", file_id)["content"]
        history = client.get("/api/conversations/" + payload["conversation_id"])
        assert history.status_code == 200
        assert f"/api/files/{file_id}/download" in history.json()["history"][-1]["content"]
        normal = client.post("/api/chat", json={"text": "¿Qué es un PDF?"})
        assert normal.status_code == 200 and normal.json()["generated_file"] is None
        assert normal.json()["answer"] == "Una respuesta normal."
        client.post("/auth/logout")
        client.post("/auth/login", json={"username": "bob", "password": "password123"})
        assert client.get(f"/api/files/{file_id}/download").status_code == 404


def test_explicit_creation_never_depends_on_model_action_json():
    class RefusingSelector:
        def answer(self, messages):
            raise RuntimeError("selector_should_not_run")
    from iachat.tool_planner import explicit_file_create_request
    from iachat.document_render import explicit_spreadsheet_columns
    assert explicit_spreadsheet_columns("Que tenga fecha, concepto, categoría, importe y un total automático.") == ["Fecha", "Concepto", "Categoría", "Importe"]
    assert explicit_spreadsheet_columns("para anotar Fecha, Concepto, Categoría e Importe, con suma automática.") == ["Fecha", "Concepto", "Categoría", "Importe"]
    for request in (
        "Creame una planilla Excel para controlar mis gastos mensuales. Que tenga fecha, concepto, categoría, importe y un total automático. Dame el archivo para descargar.",
        "Armame un Excel descargable de movimientos con una suma automática de importes.",
        "¿Podés generar un PDF descargable con un informe?",
    ):
        assert explicit_file_create_request(request)
        assert select_file_action(request, [], RefusingSelector()) == FileAction("create")
    for question in (
        "¿Cómo crear una planilla Excel?",
        "No crees un archivo Excel, explicame cómo hacerlo.",
        "¿Qué es una fórmula en Excel?",
    ):
        assert not explicit_file_create_request(question)


def test_authenticated_excel_creation_from_exact_chat_and_equivalent(monkeypatch, tmp_path):
    db = str(tmp_path / "excel-chat.sqlite")
    monkeypatch.setattr(app_module, "store", ConversationStore(db))
    monkeypatch.setattr(app_module, "file_workspace", FileWorkspace(db))
    monkeypatch.setenv("IACHAT_AUTH_SECRET", "isolated-test-secret")
    assert app_module.store.create_user("alice", "password123", admin=True)
    class ExcelBackend:
        def answer(self, messages, images=None):
            policy = messages[0]["content"] if messages[0]["role"] == "system" else ""
            if "selector de herramientas" in policy:
                raise AssertionError("explicit_creation_must_not_call_selector")
            if "generador y editor" in policy:
                return json.dumps({
                    "name": "gastos.xlsx",
                    "content": "Fecha,Concepto,Categoría,Importe\n2026-10-30,Comida,Alimentos,125.50\n",
                }, ensure_ascii=False)
            raise AssertionError("should_not_fall_back_to_normal_chat")
    monkeypatch.setattr(app_module, "get_backend", lambda: ExcelBackend())
    prompts = (
        "Creame una planilla Excel para controlar mis gastos mensuales. Que tenga fecha, concepto, categoría, importe y un total automático. Dame el archivo para descargar.",
        "Armame un Excel descargable para anotar Fecha, Concepto, Categoría e Importe, con suma automática.",
    )
    with TestClient(app_module.app, base_url="https://testserver") as client:
        assert client.post("/auth/login", json={"username": "alice", "password": "password123"}).status_code == 200
        for prompt in prompts:
            response = client.post("/api/chat", json={"text": prompt})
            assert response.status_code == 200, response.text
            data = response.json()
            info = data["generated_file"]
            assert info["name"].endswith(".xlsx")
            assert f"/api/files/{info['id']}/download" in data["answer"]
            download = client.get(f"/api/files/{info['id']}/download")
            assert download.status_code == 200 and download.content.startswith(b"PK")
            sheet = load_workbook(BytesIO(download.content)).active
            assert [sheet.cell(1, i).value for i in range(1, 5)] == [
                "Fecha", "Concepto", "Categoría", "Importe",
            ]
            assert sheet["D2"].value in (None, "")  # Ready for the user to enter a real amount.
            assert sheet["D102"].value == "=SUM(D2:D101)"
            assert sheet["D102"].data_type == "f"
            assert sheet["A102"].value == "TOTAL"


def test_explicit_file_failure_reports_error_instead_of_text_guide(monkeypatch, tmp_path):
    db = str(tmp_path / "bad-generation.sqlite")
    monkeypatch.setattr(app_module, "store", ConversationStore(db))
    monkeypatch.setattr(app_module, "file_workspace", FileWorkspace(db))
    monkeypatch.setenv("IACHAT_AUTH_SECRET", "isolated-test-secret")
    assert app_module.store.create_user("alice", "password123", admin=True)
    class BadGenerator:
        def answer(self, messages, images=None):
            return "Para crear la planilla, abrí Excel y seguí estos pasos..."
    monkeypatch.setattr(app_module, "get_backend", lambda: BadGenerator())
    with TestClient(app_module.app, base_url="https://testserver") as client:
        assert client.post("/auth/login", json={"username": "alice", "password": "password123"}).status_code == 200
        response = client.post("/api/chat", json={
            "text": "Creame una planilla Excel descargable con un total automático.",
        })
        assert response.status_code == 422
        assert response.json()["detail"] == "invalid_file_response"
        assert client.get("/api/files").json()["files"] == []
