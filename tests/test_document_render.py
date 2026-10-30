from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from fastapi.testclient import TestClient

import iachat.app as app_module
from iachat.file_workspace import FileWorkspace
from iachat.document_render import render_download
from iachat.store import ConversationStore


def test_pdf_docx_xlsx_real_bytes_and_text():
    source = "# Informe\n\n- Obra 123\n- Pendiente"
    pdf, mime = render_download("informe.pdf", source)
    assert mime == "application/pdf" and pdf.startswith(b"%PDF-")
    assert "Obra 123" in PdfReader(BytesIO(pdf)).pages[0].extract_text()
    docx, mime = render_download("informe.docx", source)
    assert mime.endswith("wordprocessingml.document")
    document = Document(BytesIO(docx))
    assert any("Obra 123" in paragraph.text for paragraph in document.paragraphs)
    excel, mime = render_download("datos.xlsx", 'nombre,valor\nObra,123\nAtaque,"=1+1"\n')
    assert mime.endswith("spreadsheetml.sheet")
    sheet = load_workbook(BytesIO(excel)).active
    assert sheet["A2"].value == "Obra" and sheet["B2"].value == "123"
    assert sheet["B3"].value == "=1+1" and sheet["B3"].data_type == "s"


def test_authenticated_downloads_and_cross_user_isolation(monkeypatch, tmp_path):
    db = str(tmp_path / "files.db")
    monkeypatch.setattr(app_module, "store", ConversationStore(db))
    monkeypatch.setattr(app_module, "file_workspace", FileWorkspace(db))
    monkeypatch.setenv("IACHAT_AUTH_SECRET", "test-only-not-production")
    assert app_module.store.create_user("alice", "password123", admin=True)
    assert app_module.store.create_user("bob", "password123")
    with TestClient(app_module.app, base_url="https://testserver") as client:
        assert client.post("/auth/login", json={"username": "alice", "password": "password123"}).status_code == 200
        for filename, content, signature in [
            ("informe.pdf", "# Informe\nTexto de prueba", b"%PDF-"),
            ("informe.docx", "# Informe\nTexto de prueba", b"PK"),
            ("datos.xlsx", "nombre,valor\nobra,123", b"PK"),
        ]:
            response = client.post("/api/files", json={"name": filename, "content": content})
            assert response.status_code == 201, response.text
            file_id = response.json()["id"]
            result = client.get(f"/api/files/{file_id}/download")
            assert result.status_code == 200 and result.content.startswith(signature)
            assert result.headers["x-content-type-options"] == "nosniff"
            assert "attachment;" in result.headers["content-disposition"]
            client.post("/auth/logout")
            client.post("/auth/login", json={"username": "bob", "password": "password123"})
            assert client.get(f"/api/files/{file_id}/download").status_code == 404
            client.post("/auth/logout")
            client.post("/auth/login", json={"username": "alice", "password": "password123"})
