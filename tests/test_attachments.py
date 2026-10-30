"""Offline attachment tests; never contact Router or consume tokens."""
import base64
import unittest
from types import SimpleNamespace
from iachat.attachments import extract, AttachmentError

def file(name, raw):
    return SimpleNamespace(name=name, content_base64=base64.b64encode(raw).decode())

class AttachmentsTest(unittest.TestCase):
    def test_text_extract(self):
        text, names = extract([file("notes.txt", b"Buenos dias")])
        self.assertIn("Buenos dias", text)
        self.assertEqual(names, ["notes.txt"])

    def test_markdown_extract(self):
        text, _ = extract([file("doc.md", b"# Titulo")])
        self.assertIn("# Titulo", text)

    def test_docx_and_xlsx_extract(self):
        import io
        from docx import Document
        from openpyxl import Workbook
        document = Document()
        document.add_heading("Informe", 1)
        document.add_paragraph("Reparación de sumidero")
        docx_data = io.BytesIO()
        document.save(docx_data)
        text, names = extract([file("informe.docx", docx_data.getvalue())])
        self.assertIn("Reparación de sumidero", text)
        self.assertEqual(names, ["informe.docx"])
        book = Workbook()
        book.active.append(["obra", "estado"])
        book.active.append(["123", "pendiente"])
        xlsx_data = io.BytesIO()
        book.save(xlsx_data)
        text, names = extract([file("planilla.xlsx", xlsx_data.getvalue())])
        self.assertIn("123 | pendiente", text)
        self.assertEqual(names, ["planilla.xlsx"])

    def test_reject_corrupt_office_file(self):
        with self.assertRaises(AttachmentError):
            extract([file("falso.docx", b"not-a-zip")])

    def test_unsupported(self):
        with self.assertRaises(AttachmentError):
            extract([file("private.exe", b"abc")])

    def test_binary_text(self):
        with self.assertRaises(AttachmentError):
            extract([file("no.txt", b"\xff\xfe")])

    def test_too_large(self):
        with self.assertRaises(AttachmentError):
            extract([file("big.txt", b"x" * 2000001)])

    def test_max_three(self):
        with self.assertRaises(AttachmentError):
            extract([file("a.txt", b"a")] * 4)

if __name__ == "__main__":
    unittest.main()
