"""Authenticated, real HTTP production smoke; isolated temporary account, no secrets in logs.

Invoked after the release has started. The privileged launcher reads the running
service environment in memory and drops to the service UID before importing the
application or touching SQLite. No existing user credentials are read or changed.
"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import pwd
import re
import secrets
import time
import sqlite3
import subprocess
import sys
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[1]
BASE = "http://127.0.0.1:8000"
QUESTIONS = (
    "¿Cuándo es el próximo partido de River Plate? Indicá fecha, horario y URL de una fuente comprobable.",
    "¿Cuál es la cotización del dólar oficial en Argentina hoy? Indicá el valor y URL de la fuente consultada.",
    "¿Cuáles son las novedades de ANSES de esta semana? Indicá una noticia reciente y URL de la fuente consultada.",
    "¿Cuándo es la próxima carrera de Fórmula 1? Indicá fecha y URL de una fuente comprobable; si no se confirma, decilo.",
    "¿Cuándo juega River Plate y cuándo es la próxima carrera de Fórmula 1? Respondé ambas con fuentes específicas.",
)
GENERIC = re.compile(
    r"(?i)(no (?:dispongo|tengo|cuento) (?:con |de )?(?:datos|informaci[oó]n) actualizad|"
    r"no (?:tengo acceso|puedo acceder) (?:a |al )?(?:internet|la web|datos en tiempo real)|"
    r"consult[aá] (?:la p[aá]gina|un sitio|sitios))"
)


def launcher() -> None:
    if os.geteuid() != 0:
        raise RuntimeError("launcher_requires_root")
    pid = int(subprocess.check_output(
        ["systemctl", "show", "iachat", "-p", "MainPID", "--value"],
        text=True, timeout=8,
    ).strip())
    if pid <= 1:
        raise RuntimeError("iachat_process_not_running")
    process = Path(f"/proc/{pid}")
    raw = (process / "environ").read_bytes()
    service_env = {}
    for entry in raw.split(b"\x00"):
        if b"=" in entry:
            key, value = entry.split(b"=", 1)
            service_env[key.decode("utf-8", "surrogateescape")] = value.decode("utf-8", "surrogateescape")
    if not service_env.get("IACHAT_AUTH_SECRET"):
        raise RuntimeError("production_auth_secret_missing")
    if service_env.get("IACHAT_BACKEND", "integrated").lower() != "integrated":
        raise RuntimeError("production_backend_not_integrated")
    uid = process.stat().st_uid
    user = pwd.getpwuid(uid)
    cwd = os.readlink(str(process / "cwd"))
    python = ROOT / ".venv/bin/python"
    if not python.is_file():
        raise RuntimeError("production_virtualenv_missing")
    service_env["PYTHONPATH"] = str(ROOT)
    service_env["IACHAT_E2E_EXPECTED_APP_SHA256"] = os.getenv("IACHAT_E2E_EXPECTED_APP_SHA256", "")
    os.chdir(cwd)
    os.initgroups(user.pw_name, user.pw_gid)
    os.setgid(user.pw_gid)
    os.setuid(uid)
    os.execve(str(python), [str(python), str(Path(__file__).resolve()), "--child"], service_env)


def cleanup_account(db_path: str, username: str) -> None:
    with sqlite3.connect(db_path, timeout=20) as connection:
        for table in ("structured_records", "structured_collections", "workspace_files", "automatic_memory_log", "memories",
                      "conversation_index", "conversations", "image_attachments"):
            if connection.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)
            ).fetchone():
                columns = {row[1] for row in connection.execute(f"PRAGMA table_info({table})")}
                if "owner" in columns:
                    connection.execute(f"DELETE FROM {table} WHERE owner=?", (username,))
        connection.execute("DELETE FROM users WHERE username=?", (username,))


def diagnose_news_feed(query: str) -> None:
    """Production-only public RSS diagnostics; print metadata, never article text.

    No credentials, model calls, user conversations, or arbitrary URLs.
    """
    import httpx
    import xml.etree.ElementTree as ET
    from iachat.web_search import (
        BING_NEWS_RSS, BING_RSS, RSS_HEADERS,
        _bing_news_publisher_url, _public_https_link,
        _relevant_direct_fallback, _clean_snippet,
    )
    feeds = (
        ("bing_news", BING_NEWS_RSS, {"q": query, "format": "rss", "setmkt": "es-AR"}),
        ("bing_general", BING_RSS, {"q": query, "format": "rss"}),
    )
    for label, endpoint, params in feeds:
        try:
            with httpx.Client(timeout=6.0, follow_redirects=False) as client:
                with client.stream("GET", endpoint, params=params, headers=RSS_HEADERS) as response:
                    status = response.status_code
                    body = bytearray()
                    for chunk in response.iter_bytes():
                        body.extend(chunk)
                        if len(body) > 400_000:
                            raise ValueError("oversized_feed")
            if status != 200:
                print("RSS_DIAG:", label, "http", status, flush=True)
                continue
            root = ET.fromstring(bytes(body))
            items = root.findall(".//channel/item")[:20]
            safe = 0
            relevant = 0
            wrappers = 0
            wrapper_schemes: set[str] = set()
            target_schemes: set[str] = set()
            for item in items:
                raw = (item.findtext("link") or "").strip()
                if "/news/apiclick.aspx?" in raw:
                    from urllib.parse import parse_qs
                    wrappers += 1
                    parsed_wrapper = urlsplit(raw)
                    wrapper_schemes.add(parsed_wrapper.scheme[:12])
                    for target in parse_qs(parsed_wrapper.query).get("url", [])[:2]:
                        target_schemes.add(urlsplit(target).scheme[:12])
                direct = _bing_news_publisher_url(raw) if label == "bing_news" else raw
                if _public_https_link(direct):
                    safe += 1
                    candidate = {
                        "title": _clean_snippet(item.findtext("title") or "", 180),
                        "snippet": _clean_snippet(item.findtext("description") or "", 450),
                    }
                    if _relevant_direct_fallback(candidate, query):
                        relevant += 1
            print("RSS_DIAG:", label, "http", status, "items", len(items),
                  "wrapped", wrappers, "wrapper_schemes", ",".join(sorted(wrapper_schemes)),
                  "target_schemes", ",".join(sorted(target_schemes)),
                  "safe_https", safe, "relevant", relevant, flush=True)
        except (httpx.HTTPError, ET.ParseError, ValueError, OverflowError, TypeError) as exc:
            print("RSS_DIAG:", label, "error", type(exc).__name__, flush=True)


def diagnose_page_failures(evidence_text: str) -> None:
    """Diagnose failed public HTTPS pages without article text, paths or secrets.

    This is run only when the normal page reader obtained zero pages. It uses
    the same validated/pinned HTTPS reader, never a bypass or a raw curl.
    """
    from iachat.web_page_reader import read_page, PageReadError
    from iachat.web_fetch import PageFetchError
    try:
        start = evidence_text.index("[")
        records, _ = json.JSONDecoder().raw_decode(evidence_text[start:])
    except (ValueError, TypeError):
        return
    if not isinstance(records, list):
        return
    attempted = 0
    for item in records:
        if not isinstance(item, dict):
            continue
        url = item.get("url")
        if not isinstance(url, str):
            continue
        try:
            host = urlsplit(url).hostname or ""
        except ValueError:
            continue
        if host == "news.google.com" or not re.fullmatch(r"[a-zA-Z0-9.-]{1,253}", host):
            continue
        if attempted >= 4:
            break
        attempted += 1
        try:
            read_page(url, timeout=3.0)
        except PageFetchError as exc:
            message = str(exc)
            status = re.search(r"HTTP_([1-5][0-9]{2})", message)
            if status:
                reason = "http_" + status.group(1)
            elif "byte limit" in message or "content length" in message:
                reason = "oversized"
            elif "compressed" in message:
                reason = "compressed"
            elif "content type" in message:
                reason = "unsupported_media"
            elif "DNS" in message:
                reason = "dns_rejected"
            elif "request failed" in message:
                reason = "transport_failure"
            else:
                reason = "safety_or_http_rejected"
        except PageReadError:
            reason = "no_readable_text"
        except ValueError:
            reason = "extraction_rejected"
        else:
            reason = "readable_on_retry"
        print("PAGE_DIAG:", "host=" + host, "reason=" + reason, flush=True)


def stage_evidence(question: str) -> tuple[int, set[str]]:
    from iachat.web_planner import plan_web_lookup_semantic
    from iachat.web_research import web_evidence_with_pages
    from iachat.backend import ProviderBackend
    plan = plan_web_lookup_semantic(question, [], ProviderBackend())
    print("PLANNER:", plan.action, plan.reason, "query=", repr(plan.query[:160]), flush=True)
    if plan.action != "search":
        raise AssertionError("planner_skipped_current_question")
    evidence = web_evidence_with_pages(plan.query, news_fallback=True)
    sources: set[str] = set()
    pages_read = 0
    try:
        start = evidence.index("[")
        records, _ = json.JSONDecoder().raw_decode(evidence[start:])
        for item in records:
            url = item.get("url")
            if isinstance(url, str) and urlsplit(url).scheme == "https":
                sources.add(url)
            if isinstance(item.get("page_excerpt_unverified"), str):
                pages_read += 1
    except (ValueError, TypeError, KeyError):
        pass
    direct_publishers = sum(urlsplit(url).hostname != "news.google.com" for url in sources)
    print("WEB_EVIDENCE:", "sources", len(sources),
          "direct_publishers", direct_publishers, "pages_read", pages_read,
          "bytes", len(evidence.encode("utf-8")), flush=True)
    print("WEB_SOURCE_HOSTS:", ",".join(sorted({urlsplit(u).hostname or "" for u in sources})[:5]), flush=True)
    if not pages_read:
        diagnose_page_failures(evidence)
        diagnose_news_feed(plan.query)
    return len(sources), sources


# Live free-provider checks should not consume a minute's shared quota in
# a burst. Pace actual authenticated HTTP calls; keep all assertions and
# never retry a mutation that may have committed before an error response.
_last_model_response_at = 0.0


def paced_model_post(session, path, **kwargs):
    global _last_model_response_at
    if _last_model_response_at:
        delay = max(0.0, 20.0 - (time.monotonic() - _last_model_response_at))
        if delay:
            print("E2E_FREE_PROVIDER_PACING: waiting", round(delay, 1), "s", flush=True)
            time.sleep(delay)
    try:
        return session.post(path, **kwargs)
    finally:
        _last_model_response_at = time.monotonic()


def child() -> None:
    import httpx
    from iachat.store import ConversationStore
    from iachat.backend import ProviderBackend

    # All live model requests go through the existing fail-closed free-model allowlist.
    if not ProviderBackend.configured():
        raise AssertionError("free_provider_not_configured")
    # Model IDs and configured flags only; never print keys or provider responses.
    for provider_name, env_key, _, defaults in ProviderBackend.PROVIDERS:
        eligible = ProviderBackend._eligible_models(provider_name, defaults)
        print("FREE_PROVIDER:", provider_name,
              "configured=" + str(bool(os.getenv(env_key, "").strip())),
              "models=" + ",".join(eligible), flush=True)
    db_path = os.getenv("IACHAT_DB", "iachat.db")
    if db_path == ":memory:":
        raise AssertionError("production_database_is_memory_only")
    db_path = str(Path(db_path).resolve())
    if not Path(db_path).is_file():
        raise AssertionError("production_database_not_found")
    expected = os.getenv("IACHAT_E2E_EXPECTED_APP_SHA256", "")
    actual = hashlib.sha256((ROOT / "iachat/app.py").read_bytes()).hexdigest()
    print("RELEASE_APP_SHA256:", actual, flush=True)
    if expected and actual != expected:
        raise AssertionError("deployed_source_hash_mismatch")

    username = "e2e-" + secrets.token_hex(8)
    password = secrets.token_urlsafe(32)
    other_username = "e2e-" + secrets.token_hex(8)
    other_password = secrets.token_urlsafe(32)
    store = ConversationStore(db_path)
    if not store.create_user(username, password):
        raise AssertionError("temporary_user_creation_failed")
    if not store.create_user(other_username, other_password):
        cleanup_account(db_path, username)
        raise AssertionError("second_temporary_user_creation_failed")
    with sqlite3.connect(db_path) as audit_db:
        print("PRODUCTION_ACCOUNT_COUNT:", audit_db.execute("SELECT count(*) FROM users").fetchone()[0], flush=True)
    try:
        with httpx.Client(base_url=BASE, timeout=160.0, follow_redirects=False) as client:
            login = client.post("/auth/login", json={"username": username, "password": password})
            print("AUTH_LOGIN_HTTP:", login.status_code, flush=True)
            if login.status_code != 200:
                raise AssertionError("real_login_failed")
            cookie = login.cookies.get("iachat_session")
            if not cookie:
                raise AssertionError("login_cookie_missing")
            # The production cookie has Secure; the internal HTTP loopback smoke
            # must explicitly send it (the public browser continues using HTTPS).
            client.headers["Cookie"] = "iachat_session=" + cookie
            status = client.get("/auth/status")
            if status.status_code != 200 or status.json().get("username") != username:
                raise AssertionError("real_cookie_auth_failed")
            spec = client.get("/openapi.json")
            if spec.status_code != 200 or "/api/files/generate" not in spec.json().get("paths", {}):
                raise AssertionError("new_release_not_served")
            print("AUTH_COOKIE_AND_FILE_ROUTES: OK", flush=True)

            failures = []
            # The actual user-observed regression: a conversational Google
            # account probe must never invoke web evidence or show RSS sources.
            # One authenticated HTTP call uses a real allowlisted free model.
            try:
                probe_started = time.monotonic()
                probe = paced_model_post(client, "/api/chat", json={
                    "text": "Probando si anda IAchat con mi cuenta de mail",
                })
                probe_data = probe.json() if probe.status_code == 200 else {}
                probe_answer = probe_data.get("answer", "")
                probe_sources = probe_data.get("web_sources", [])
                probe_ok = (
                    probe.status_code == 200
                    and isinstance(probe_answer, str) and bool(probe_answer.strip())
                    and not GENERIC.search(probe_answer)
                    and probe_sources == []
                    and "Fuentes recuperadas" not in probe_answer
                    and "https://" not in probe_answer
                )
                print(
                    "GOOGLE_BANAL_CHAT_NO_WEB:",
                    "OK" if probe_ok else f"HTTP_{probe.status_code}",
                    "sources=" + str(len(probe_sources) if isinstance(probe_sources, list) else -1),
                    "elapsed_ms=" + str(int((time.monotonic() - probe_started) * 1000)),
                    flush=True,
                )
                if not probe_ok:
                    failures.append("google_banal_chat_web_or_answer_regression")
                if probe_ok:
                    restored = client.get("/api/conversations/" + probe_data["conversation_id"])
                    restored_last = (
                        restored.json().get("history", [{}])[-1]
                        if restored.status_code == 200 else {}
                    )
                    persisted_ok = (
                        restored_last.get("web_sources") == []
                        and isinstance(restored_last.get("elapsed_ms"), int)
                        and isinstance(restored_last.get("completed_at"), str)
                    )
                    print("GOOGLE_BANAL_RELOAD_METADATA:",
                          "OK" if persisted_ok else "FAIL", flush=True)
                    if not persisted_ok:
                        failures.append("google_banal_reload_metadata_missing")
            except (httpx.HTTPError, ValueError, TypeError) as exc:
                print("GOOGLE_BANAL_CHAT_ERROR:", type(exc).__name__, flush=True)
                failures.append("google_banal_chat_http_error")
            # Ordinary inference must work independently of web search.
            try:
                normal = paced_model_post(client, "/api/chat", json={"text": "¿Cuánto es 7 por 8? Respondé con el resultado."})
                normal_answer = normal.json().get("answer", "") if normal.status_code == 200 else ""
                normal_ok = normal.status_code == 200 and "56" in normal_answer
                print("NORMAL_CHAT:", "OK" if normal_ok else f"HTTP_{normal.status_code}", flush=True)
                if not normal_ok:
                    failures.append("normal_chat_failed")
            except (httpx.HTTPError, ValueError, TypeError) as exc:
                print("NORMAL_CHAT_ERROR:", type(exc).__name__, flush=True)
                failures.append("normal_chat_http_error")
            from iachat.web_planner import split_independent_search_questions
            for idx, question in enumerate(QUESTIONS, 1):
                # A compound question must search each topic independently.
                # This is read-only public evidence; no user data or secrets.
                for part_no, subquestion in enumerate(split_independent_search_questions(question), 1):
                    count, _ = stage_evidence(subquestion)
                    if not count:
                        # Independent RSS preflight is inherently non-atomic:
                        # the actual chat performs its own fresh lookup. Gate
                        # on its returned sources and topic coverage below,
                        # not on an unrelated earlier feed response.
                        print(f"WEB_PREFLIGHT_{idx}_{part_no}: no_sources_on_separate_lookup", flush=True)
                try:
                    response = paced_model_post(client, "/api/chat", json={"text": question})
                    data = response.json() if response.status_code == 200 else {}
                    answer = data.get("answer", "")
                    # The actual chat response carries the exact consulted links;
                    # an independent second RSS request can return different URLs.
                    returned = data.get("web_sources", [])
                    if response.status_code == 200 and data.get("conversation_id"):
                        restored = client.get("/api/conversations/" + data["conversation_id"])
                        restored_last = (
                            restored.json().get("history", [{}])[-1]
                            if restored.status_code == 200 else {}
                        )
                        persisted_ok = (
                            restored_last.get("web_sources") == returned
                            and isinstance(restored_last.get("elapsed_ms"), int)
                            and isinstance(restored_last.get("completed_at"), str)
                        )
                        print(f"CHAT_{idx}_RELOAD_METADATA:",
                              "OK" if persisted_ok else "FAIL", flush=True)
                        if not persisted_ok:
                            failures.append(f"chat_{idx}_reload_metadata_missing")
                    source_urls = {
                        item["url"] for item in returned
                        if isinstance(item, dict) and isinstance(item.get("url"), str)
                        and urlsplit(item["url"]).scheme == "https"
                    } if isinstance(returned, list) else set()
                    urls = set(re.findall(r"https://[^\s)<>]+", answer))
                    ordered_sources = [
                        item for item in returned if isinstance(item, dict)
                        and isinstance(item.get("url"), str)
                    ] if isinstance(returned, list) else []
                    direct_count = sum(
                        urlsplit(item["url"]).hostname != "news.google.com"
                        for item in ordered_sources
                    )
                    pages_read = sum(
                        item.get("evidence_type") == "page_read"
                        for item in ordered_sources
                    )
                    first_three_direct = sum(
                        urlsplit(item["url"]).hostname != "news.google.com"
                        for item in ordered_sources[:3]
                    )
                    print(f"CHAT_{idx}_PROVENANCE: direct={direct_count} "
                          f"pages_read={pages_read} first_three_direct={first_three_direct}",
                          flush=True)
                    if direct_count and not first_three_direct:
                        failures.append(f"chat_{idx}_publisher_links_hidden")
                    matched = any(u.rstrip(".,;") in source_urls for u in urls)
                    invalid_citations = [u for u in urls if u.rstrip(".,;") not in source_urls]
                    generic = bool(GENERIC.search(answer))
                    cited_hosts = sorted({
                        urlsplit(u).hostname or "" for u in urls
                        if urlsplit(u).scheme == "https"
                    })[:5]
                    print(f"CHAT_{idx}: http={response.status_code} chars={len(answer)} "
                          f"generic={generic} model_urls={len(urls)} "
                          f"cited_hosts={','.join(cited_hosts)} actual_sources={len(source_urls)} "
                          f"model_cited_actual={matched}", flush=True)
                    # The UI shows these exact sources, marked unverified.
                    # Do not require the model to reproduce their long RSS URLs.
                    if response.status_code != 200 or generic or len(answer.strip()) < 35 or not source_urls:
                        failures.append(f"chat_{idx}_missing_grounded_answer")
                    if invalid_citations:
                        failures.append(f"chat_{idx}_cites_unconsulted_url")
                    if idx in (4, 5) and response.status_code == 200:
                        from iachat.fresh_event_guard import reject_stale_next_event_answer
                        if reject_stale_next_event_answer(question, answer) != answer:
                            failures.append(f"chat_{idx}_stale_upcoming_date")
                        # A suppressed date is safer than misinformation, but
                        # it is not proof that the next event was verified.
                        suppressed = "La información obtenida contiene una fecha que ya pasó" in answer
                        print(f"CHAT_{idx}_UPCOMING_DATE_SUPPRESSED: {suppressed}", flush=True)
                        if idx == 5:
                            both_topics = (
                                bool(re.search(r"(?i)\briver\b", answer))
                                and bool(re.search(r"(?i)f[oó]rmula\s*1|\bf1\b|carrera|gran premio", answer))
                            )
                            print("CHAT_5_BOTH_TOPICS:", both_topics, flush=True)
                            if not both_topics and not suppressed:
                                failures.append("chat_5_missing_one_topic")
                            # Topic coverage in returned source titles is a
                            # necessary sanity check, NOT factual corroboration.
                            source_titles = " ".join(str(item.get("title", "")) for item in ordered_sources)
                            river_source = bool(re.search(r"(?i)\briver\b", source_titles))
                            racing_source = bool(re.search(
                                r"(?i)f[oó]rmula\s*1|\bf1\b|gran premio|\bgp\b|carrera",
                                source_titles,
                            ))
                            print("CHAT_5_SOURCE_TOPIC_COVERAGE:",
                                  "river=" + str(river_source),
                                  "racing=" + str(racing_source), flush=True)
                            if not (river_source and racing_source):
                                failures.append("chat_5_sources_do_not_cover_both_topics")
                    direct_hosts = {
                        urlsplit(item["url"]).hostname for item in ordered_sources
                        if urlsplit(item["url"]).hostname != "news.google.com"
                    }
                    rss_publishers = {
                        item["rss_publisher_unverified"].casefold()
                        for item in ordered_sources
                        if isinstance(item.get("rss_publisher_unverified"), str)
                        and item["rss_publisher_unverified"].strip()
                    }
                    print(f"CHAT_{idx}_RSS_PUBLISHERS: {len(rss_publishers)}", flush=True)
                    # Multiple named publishers are an explicitly UNVERIFIED
                    # RSS fallback, not a substitute for corroborating facts.
                    if pages_read == 0 and len(direct_hosts) < 2 and len(rss_publishers) < 2:
                        failures.append(f"chat_{idx}_insufficient_independent_rss_hosts")
                except (httpx.HTTPError, ValueError, TypeError) as exc:
                    print(f"CHAT_{idx}: error={type(exc).__name__}", flush=True)
                    failures.append(f"chat_{idx}_http_error")

            # Deterministic real authenticated file creation and byte-exact download.
            payload = {"name": "e2e-verificacion.csv", "content": "nombre,valor\nprueba,123\n"}
            saved = client.post("/api/files", json=payload)
            if saved.status_code != 201:
                failures.append("file_create_failed")
            else:
                file_id = saved.json()["id"]
                downloaded = client.get(f"/api/files/{file_id}/download")
                match = downloaded.status_code == 200 and downloaded.content == payload["content"].encode()
                print("FILE_DOWNLOAD:", "OK" if match else "FAILED", flush=True)
                if not match:
                    failures.append("file_download_mismatch")
            # Verify actual binary office files over authenticated HTTP, not
            # just a health check or the Python renderer in isolation.
            from io import BytesIO
            from docx import Document
            from openpyxl import load_workbook
            from pypdf import PdfReader
            office_cases = [
                ("e2e-documento.pdf", "# Informe E2E\\nTexto de prueba"),
                ("e2e-documento.docx", "# Informe E2E\\nTexto de prueba"),
                ("e2e-datos.xlsx", "nombre,valor\\nprueba,123\\nformula,=1+1"),
            ]
            for filename, raw_content in office_cases:
                raw_content = raw_content.replace("\\n", "\n")
                response = client.post("/api/files", json={"name": filename, "content": raw_content})
                ok = response.status_code == 201
                if ok:
                    download = client.get(f"/api/files/{response.json()['id']}/download")
                    ok = download.status_code == 200
                    if ok:
                        try:
                            if filename.endswith(".pdf"):
                                ok = download.content.startswith(b"%PDF-") and "Informe E2E" in PdfReader(BytesIO(download.content)).pages[0].extract_text()
                            elif filename.endswith(".docx"):
                                ok = any("Informe E2E" in p.text for p in Document(BytesIO(download.content)).paragraphs)
                            else:
                                sheet = load_workbook(BytesIO(download.content)).active
                                ok = sheet["B2"].value == 123 and sheet["B3"].value == "=1+1" and sheet["B3"].data_type == "s"
                        except Exception:
                            ok = False
                print("OFFICE_DOWNLOAD:", filename.rsplit(".", 1)[-1], "OK" if ok else "FAILED", flush=True)
                if not ok:
                    failures.append("office_" + filename.rsplit(".", 1)[-1] + "_failed")
            # Real HTTP document analysis: one inference with DOCX and XLSX.
            # The markers are synthetic and never disclose production data.
            try:
                import base64
                from io import BytesIO
                from docx import Document
                from openpyxl import Workbook
                doc = Document()
                doc.add_paragraph("Código del documento: LUCERO-482")
                doc_buffer = BytesIO()
                doc.save(doc_buffer)
                book = Workbook()
                book.active.append(["concepto", "codigo"])
                book.active.append(["referencia", "MATE-731"])
                xlsx_buffer = BytesIO()
                book.save(xlsx_buffer)
                attachment_response = paced_model_post(client, "/api/chat", json={
                    "text": "Leé los dos documentos adjuntos y decime los dos códigos que aparecen, sin inventar otros.",
                    "attachments": [
                        {"name": "prueba.docx", "content_base64": base64.b64encode(doc_buffer.getvalue()).decode()},
                        {"name": "prueba.xlsx", "content_base64": base64.b64encode(xlsx_buffer.getvalue()).decode()},
                    ],
                })
                attachment_answer = attachment_response.json().get("answer", "") if attachment_response.status_code == 200 else ""
                attachment_ok = (
                    attachment_response.status_code == 200
                    and "LUCERO-482" in attachment_answer.upper()
                    and "MATE-731" in attachment_answer.upper()
                )
                print("DOCX_XLSX_CHAT_ANALYSIS:", "OK" if attachment_ok else f"HTTP_{attachment_response.status_code}", flush=True)
                if not attachment_ok:
                    failures.append("docx_xlsx_chat_analysis_failed")
            except (httpx.HTTPError, ValueError, TypeError) as exc:
                print("DOCX_XLSX_CHAT_ANALYSIS_ERROR:", type(exc).__name__, flush=True)
                failures.append("docx_xlsx_chat_analysis_error")
            # Model-generated file: real authenticated provider inference.
            try:
                generated = paced_model_post(client, "/api/files/generate", json={
                    "instruction": "Creá un archivo llamado resumen-e2e.txt cuyo contenido sea exactamente: Prueba IAchat OK."
                })
                generated_ok = generated.status_code == 201
                if generated_ok:
                    info = generated.json()
                    result = client.get(f"/api/files/{info['id']}/download")
                    generated_ok = result.status_code == 200 and b"Prueba IAchat OK" in result.content
                print("MODEL_FILE_GENERATION:", "OK" if generated_ok else f"HTTP_{generated.status_code}", flush=True)
                if not generated_ok:
                    failures.append("model_file_generation_failed")
            except (httpx.HTTPError, ValueError, TypeError) as exc:
                print("MODEL_FILE_GENERATION_ERROR:", type(exc).__name__, flush=True)
                failures.append("model_file_generation_error")
            # Regression for the actual 30/10 incident. These are real
            # authenticated conversations using the configured FREE provider,
            # followed by HTTP downloads and OOXML inspection. Fail deployment
            # if either prompt yields a guide instead of a genuine XLSX.
            excel_requests = (
                "Creame una planilla Excel para controlar mis gastos mensuales. Que tenga fecha, concepto, categoría, importe y un total automático. Dame el archivo para descargar.",
                "Armame un Excel descargable para anotar Fecha, Concepto, Categoría e Importe, con suma automática.",
            )
            for index, instruction in enumerate(excel_requests, 1):
                ok = False
                status_code = 0
                try:
                    response = paced_model_post(client, "/api/chat", json={"text": instruction})
                    status_code = response.status_code
                    if response.status_code == 200:
                        payload = response.json()
                        info = payload.get("generated_file")
                        if (isinstance(info, dict)
                                and isinstance(info.get("id"), str)
                                and str(info.get("name", "")).lower().endswith(".xlsx")
                                and f"/api/files/{info['id']}/download" in payload.get("answer", "")):
                            download = client.get(f"/api/files/{info['id']}/download")
                            if download.status_code == 200 and download.content.startswith(b"PK"):
                                sheet = load_workbook(BytesIO(download.content)).active
                                headers = [sheet.cell(1, i).value for i in range(1, 5)]
                                total_row = sheet.max_row
                                formula = sheet.cell(total_row, 4)
                                ok = (
                                    headers == ["Fecha", "Concepto", "Categoría", "Importe"]
                                    and sheet.cell(total_row, 1).value == "TOTAL"
                                    and formula.data_type == "f"
                                    and bool(re.fullmatch(r"=SUM\(D2:D[0-9]{2,4}\)", str(formula.value)))
                                    and sheet.parent.calculation.calcMode == "auto"
                                )
                except (httpx.HTTPError, ValueError, TypeError, KeyError, OSError):
                    ok = False
                print(f"CHAT_EXCEL_AUTOTOTAL_{index}:",
                      "OK" if ok else f"FAILED_HTTP_{status_code}", flush=True)
                if not ok:
                    failures.append(f"chat_excel_autototal_{index}_failed")
            # One complete authenticated LIVE free-model sequence for persistent
            # structured records, followed by another browser session, a second
            # user and a non-financial collection. Only synthetic test accounts.
            try:
                def data_chat(session, prompt):
                    response = paced_model_post(session, "/api/chat", json={"text": prompt})
                    if response.status_code != 200:
                        raise AssertionError("structured_chat_http_" + str(response.status_code))
                    return response.json()

                first = data_chat(client, "creá planilla gastos fecha concepto categoría importe")
                if "colección" not in first["answer"].casefold():
                    raise AssertionError("structured_create_failed")
                second = data_chat(client, "anotá 40000 nafta hoy")
                if "40000" not in second["answer"]:
                    raise AssertionError("structured_add_failed")
                summed = data_chat(client, "¿cuánto gasté en transporte?")
                if "40000" not in summed["answer"]:
                    raise AssertionError("structured_sum_failed")
                corrected = data_chat(client, "corregí el importe a 45000")
                if "45000" not in corrected["answer"]:
                    raise AssertionError("structured_update_failed")
                with httpx.Client(base_url=BASE, timeout=160.0, follow_redirects=False) as another_device:
                    logged = another_device.post("/auth/login", json={"username": username, "password": password})
                    if logged.status_code != 200:
                        raise AssertionError("same_user_second_session_login_failed")
                    another_device.headers["Cookie"] = "iachat_session=" + logged.cookies["iachat_session"]
                    same_history = another_device.get("/api/conversations/" + first["conversation_id"])
                    if same_history.status_code != 200:
                        raise AssertionError("same_user_history_not_synchronized")
                    retained = data_chat(another_device, "¿cuánto gasté en transporte?")
                    if "45000" not in retained["answer"]:
                        raise AssertionError("structured_cross_session_persistence_failed")
                    export = data_chat(another_device, "descargá Excel")
                    info = export.get("generated_file")
                    if not isinstance(info, dict) or not str(info.get("name", "")).endswith(".xlsx"):
                        raise AssertionError("structured_export_missing_file")
                    file_id = info["id"]
                    if f"/api/files/{file_id}/download" not in export["answer"]:
                        raise AssertionError("structured_export_missing_link")
                    download = another_device.get(f"/api/files/{file_id}/download")
                    if download.status_code != 200 or not download.content.startswith(b"PK"):
                        raise AssertionError("structured_export_download_failed")
                    book = load_workbook(BytesIO(download.content))
                    sheet = book.active
                    headers = [sheet.cell(1, i).value for i in range(1, 5)]
                    if headers != ["Fecha", "Concepto", "Categoría", "Importe"]:
                        raise AssertionError("structured_export_headers_failed")
                    if sheet["D2"].value != 45000 or sheet["D3"].value != "=SUM(D2:D2)":
                        raise AssertionError("structured_export_values_or_formula_failed")
                    if sheet["D3"].data_type != "f":
                        raise AssertionError("structured_export_formula_not_real")
                    # New conversation and session share the same stored data.
                    if retained["conversation_id"] == corrected["conversation_id"]:
                        raise AssertionError("structured_new_conversation_not_created")
                print("STRUCTURED_LIVE_EXPENSES_CREATE_ADD_SUM_UPDATE_EXPORT: OK", flush=True)
                print("STRUCTURED_SAME_USER_SECOND_SESSION: OK", flush=True)
                with httpx.Client(base_url=BASE, timeout=160.0, follow_redirects=False) as other:
                    login_other = other.post("/auth/login", json={
                        "username": other_username, "password": other_password
                    })
                    if login_other.status_code != 200:
                        raise AssertionError("other_user_login_failed")
                    other.headers["Cookie"] = "iachat_session=" + login_other.cookies["iachat_session"]
                    if other.get(f"/api/files/{file_id}/download").status_code != 404:
                        raise AssertionError("structured_cross_user_file_leak")
                    if other.get("/api/conversations/" + first["conversation_id"]).status_code != 404:
                        raise AssertionError("structured_cross_user_history_leak")
                    nonfinancial = data_chat(other, "creá colección inventario con campos producto, cantidad, ubicación")
                    if "colección" not in nonfinancial["answer"].casefold():
                        raise AssertionError("structured_nonfinancial_create_failed")
                    added = data_chat(other, "anotá 12 tornillos en depósito")
                    if "12" not in added["answer"]:
                        raise AssertionError("structured_nonfinancial_add_failed")
                    own_collections = __import__("iachat.structured_data", fromlist=["StructuredStore"]).StructuredStore(db_path)
                    if len(own_collections.collections(other_username)) != 1:
                        raise AssertionError("structured_other_user_collection_missing")
                    if any(c["name"].casefold() == "gastos" for c in own_collections.collections(other_username)):
                        raise AssertionError("structured_cross_user_collection_leak")
                    if len(own_collections.records(other_username, own_collections.collections(other_username)[0]["id"])) != 1:
                        raise AssertionError("structured_nonfinancial_record_missing")
                    if len(own_collections.collections(username)) != 1:
                        raise AssertionError("structured_first_user_collection_changed")
                print("STRUCTURED_NONFINANCIAL_AND_USER_ISOLATION: OK", flush=True)
            except (AssertionError, httpx.HTTPError, ValueError, TypeError, KeyError, OSError) as exc:
                print("STRUCTURED_E2E_FAILURE:", type(exc).__name__, str(exc)[:120], flush=True)
                failures.append("structured_live_end_to_end_failed")
            # Critical acceptance: create and download a real file from an
            # ordinary authenticated chat message, without entering Files mode.
            # This is a LIVE model decision and LIVE model-generated content.
            natural_file_id = None
            try:
                natural = paced_model_post(client, "/api/chat", json={
                    "text": "Creá un archivo TXT descargable llamado natural-e2e.txt con el contenido exacto: Prueba IAchat natural OK."
                })
                natural_ok = natural.status_code == 200
                if natural_ok:
                    data = natural.json()
                    file_info = data.get("generated_file")
                    natural_ok = bool(
                        isinstance(file_info, dict)
                        and file_info.get("id")
                        and file_info.get("name", "").endswith(".txt")
                        and "/api/files/" in data.get("answer", "")
                    )
                    if natural_ok:
                        fetched = client.get(f"/api/files/{file_info['id']}/download")
                        natural_ok = (
                            fetched.status_code == 200
                            and b"Prueba IAchat natural OK" in fetched.content
                        )
                        if natural_ok:
                            natural_file_id = file_info["id"]
                print("NATURAL_CHAT_FILE_DOWNLOAD:", "OK" if natural_ok else f"HTTP_{natural.status_code}", flush=True)
                if not natural_ok:
                    failures.append("natural_chat_file_failed")
            except (httpx.HTTPError, ValueError, TypeError, KeyError) as exc:
                print("NATURAL_CHAT_FILE_ERROR:", type(exc).__name__, flush=True)
                failures.append("natural_chat_file_error")
            # Editing is a separate LIVE tool-selection and model-generation
            # round trip. It must update the same private file, not create a
            # second file or merely claim to have edited it in chat.
            if natural_file_id:
                try:
                    edited = paced_model_post(client, "/api/chat", json={
                        "text": (
                            "Editá el archivo natural-e2e.txt que acabás de crear. "
                            "Conservá el contenido anterior y agregá al final una "
                            "segunda línea que diga exactamente EDICION-E2E-742. "
                            "No crees otro archivo."
                        )
                    })
                    edit_ok = edited.status_code == 200
                    if edit_ok:
                        edited_info = edited.json().get("generated_file")
                        edit_ok = (
                            isinstance(edited_info, dict)
                            and edited_info.get("id") == natural_file_id
                        )
                        if edit_ok:
                            downloaded_edit = client.get(
                                f"/api/files/{natural_file_id}/download"
                            )
                            edit_ok = (
                                downloaded_edit.status_code == 200
                                and b"Prueba IAchat natural OK" in downloaded_edit.content
                                and b"EDICION-E2E-742" in downloaded_edit.content
                            )
                    print("NATURAL_CHAT_FILE_EDIT:", "OK" if edit_ok else f"HTTP_{edited.status_code}", flush=True)
                    if not edit_ok:
                        failures.append("natural_chat_file_edit_failed")
                except (httpx.HTTPError, ValueError, TypeError, KeyError) as exc:
                    print("NATURAL_CHAT_FILE_EDIT_ERROR:", type(exc).__name__, flush=True)
                    failures.append("natural_chat_file_edit_error")
            else:
                print("NATURAL_CHAT_FILE_EDIT: SKIPPED_CREATE_FAILED", flush=True)
            if failures:
                raise AssertionError("E2E_FAILED:" + ",".join(failures))
            print("E2E_AUTH_WEB_CHAT_FILES: PASS", flush=True)
    finally:
        cleanup_account(db_path, username)
        cleanup_account(db_path, other_username)
        print("TEMPORARY_ACCOUNTS_AND_DATA: CLEANED", flush=True)


if __name__ == "__main__":
    try:
        if len(sys.argv) != 2:
            raise RuntimeError("usage: script --launch|--child")
        if sys.argv[1] == "--launch":
            launcher()
        elif sys.argv[1] == "--child":
            child()
        else:
            raise RuntimeError("invalid_mode")
    except Exception as exc:
        # No secrets, response bodies or production data in logs.
        print("E2E_FAILURE:", type(exc).__name__, str(exc)[:200], flush=True)
        sys.exit(1)
